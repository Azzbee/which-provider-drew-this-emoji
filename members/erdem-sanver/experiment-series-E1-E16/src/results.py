"""Experiment records: one JSON per experiment -> regenerated Excel (the course template).

The Excel is never edited by hand: ``python -m src.results`` rebuilds it from experiments/*.json.
"""
import json
from dataclasses import asdict, dataclass, field
from datetime import date

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from sklearn.metrics import accuracy_score, precision_score

from src.config import AUTHOR, EXPERIMENTS_DIR, RESULTS_XLSX

# Column order of the course template "Strategies Computer Vision Project.xlsx" (+ our 2 extra columns)
COLUMNS = ["Name", "Model Tried", "Method Tried", "Number of epochs trained", "Date",
           "Train Values Accuracy", "Train Value Precision", "Ratio train", "Ratio Test",
           "Test Accuracy", "Test Precision", "Explanation", "Data preprocessing", "Explain",
           "Kaggle Public Score", "Experiment ID"]


@dataclass
class ExperimentRecord:
    exp_id: str                 # e.g. "E2_simple_cnn"
    model: str
    method: str
    epochs: int | None          # None for non-iterative models (SVM, ...)
    train_acc: float | None     # None for ensembles (no single training set fit)
    train_prec: float | None
    val_acc: float | None       # None for --full runs (no held-out data)
    val_prec: float | None
    explanation: str            # result + interpretation, written after looking at the numbers
    preprocessing: str
    preprocessing_why: str
    ratio_train: float = 0.8
    ratio_test: float = 0.2
    kaggle_public: float | None = None
    name: str = AUTHOR
    date: str = field(default_factory=lambda: date.today().isoformat())
    extra: dict = field(default_factory=dict)  # seed, img size, runtime, confusion matrix, ...

    def save(self) -> None:
        EXPERIMENTS_DIR.mkdir(exist_ok=True)
        (EXPERIMENTS_DIR / f"{self.exp_id}.json").write_text(json.dumps(asdict(self), indent=2))


def scores(y_true: np.ndarray, y_pred: np.ndarray) -> tuple[float, float]:
    """(accuracy, macro precision) — macro so the small classes (mozilla 4%) weigh equally."""
    return (float(accuracy_score(y_true, y_pred)),
            float(precision_score(y_true, y_pred, average="macro", zero_division=0)))


def _r(x: float | None) -> float | None:
    return None if x is None else round(x, 4)


def load_records() -> list[ExperimentRecord]:
    recs = [ExperimentRecord(**json.loads(p.read_text())) for p in EXPERIMENTS_DIR.glob("*.json")]
    return sorted(recs, key=lambda r: (r.date, r.exp_id))


def export_excel() -> int:
    """Rebuild the results workbook from all experiment JSONs. Returns number of rows written."""
    recs = load_records()
    wb = Workbook()
    ws = wb.active
    ws.title = "Feuille 1"
    ws.append(COLUMNS)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in recs:
        ws.append([r.name, r.model, r.method, r.epochs, r.date,
                   _r(r.train_acc), _r(r.train_prec), r.ratio_train, r.ratio_test,
                   _r(r.val_acc), _r(r.val_prec), r.explanation, r.preprocessing,
                   r.preprocessing_why, r.kaggle_public, r.exp_id])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for col, width in zip("ABCDEFGHIJKLMNOP", [14, 22, 30, 10, 11, 10, 10, 8, 8, 10, 10, 60, 40, 50, 10, 22]):
        ws.column_dimensions[col].width = width
    RESULTS_XLSX.parent.mkdir(exist_ok=True)
    wb.save(RESULTS_XLSX)
    return len(recs)


if __name__ == "__main__":
    print(f"{export_excel()} experiments -> {RESULTS_XLSX}")
